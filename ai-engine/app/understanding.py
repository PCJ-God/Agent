"""问题理解层（格式归一化）：把任意输入转成「规范文字 + 元数据」。

- 纯文本（md/txt/json/csv/log）：原样解码
- docx：python-docx 抽段落 + 表格文字
- pdf：pymupdf 抽文字层；扫描版转图片走视觉理解
- 图片（png/jpg/jpeg/webp）：强视觉模型生成「结构化描述 + 文字」（VISION_MODEL，走 OpenAI 兼容接口）
解析后文字超 MAX_TEXT_CHARS 抛 ValueError（对应「文档内容过长」）。
"""
from __future__ import annotations

import base64
from io import BytesIO

import fitz  # pymupdf
import httpx

from app.config import DASHSCOPE_API_KEY, DASHSCOPE_CHAT_URL, MAX_TEXT_CHARS, VISION_MODEL

TEXT_EXTS = (".md", ".txt", ".json", ".csv", ".log")
IMAGE_EXTS = (".png", ".jpg", ".jpeg", ".webp")

# 视觉模型提示词：产出「图片描述 + 文字内容」两部分，便于注入对话上下文
VISION_PROMPT = (
    "请阅读这张图片并用中文输出两部分：\n"
    "1）图片描述：说明图片类型与主要内容；若是图表、表格、流程图或示意图，"
    "请结构化描述其标题、坐标轴/表头、数据、趋势与结构关系。\n"
    "2）文字内容：逐字提取图中所有可见文字，尽量保持原有顺序与结构。\n"
    "若某部分没有内容，写「无」。"
)


def parse(filename: str, content: bytes) -> tuple[str, dict]:
    """解析文档为纯文本。返回 (文字, 元数据)；超长/不支持抛 ValueError。"""
    lower = filename.lower()
    if lower.endswith(TEXT_EXTS):
        text = content.decode("utf-8", errors="ignore")
        meta = {"filename": filename, "type": "text"}
    elif lower.endswith(".docx"):
        text = _parse_docx(content)
        meta = {"filename": filename, "type": "docx"}
    elif lower.endswith(".pdf"):
        text = _parse_pdf(content)
        meta = {"filename": filename, "type": "pdf"}
    elif lower.endswith(IMAGE_EXTS):
        text = _describe_image(content)
        meta = {"filename": filename, "type": "image"}
    else:
        raise ValueError(f"不支持的文件类型: {lower}")

    text = text.strip()
    if len(text) > MAX_TEXT_CHARS:
        raise ValueError(f"文档内容过长（超过 {MAX_TEXT_CHARS} 字符），请拆分或精简后重试")
    return text, meta


def _parse_docx(content: bytes) -> str:
    import docx

    doc = docx.Document(BytesIO(content))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _parse_pdf(content: bytes) -> str:
    doc = fitz.open(stream=content, filetype="pdf")
    parts = [page.get_text() for page in doc]
    text = "\n".join(parts).strip()
    if text:
        return text
    # 扫描版 PDF：逐页渲染成图片走视觉理解
    return "\n".join(
        _describe_image(page.get_pixmap(dpi=150).tobytes("png")) for page in doc
    )


def _img_mime(content: bytes) -> str:
    if content[:8] == b"\x89PNG\r\n\x1a\n":
        return "image/png"
    if content[:3] == b"\xff\xd8\xff":
        return "image/jpeg"
    if content[:4] == b"RIFF" and content[8:12] == b"WEBP":
        return "image/webp"
    return "image/png"


def _describe_image(content: bytes) -> str:
    """用强视觉模型对图片生成「结构化描述 + 文字」，作为归一化文本返回。"""
    b64 = base64.b64encode(content).decode()
    data_uri = f"data:{_img_mime(content)};base64,{b64}"
    payload = {
        "model": VISION_MODEL,
        "messages": [
            {
                "role": "user",
                "content": [
                    {"type": "image_url", "image_url": {"url": data_uri}},
                    {"type": "text", "text": VISION_PROMPT},
                ],
            }
        ],
    }
    with httpx.Client(timeout=120) as client:
        resp = client.post(
            DASHSCOPE_CHAT_URL,
            headers={
                "Authorization": f"Bearer {DASHSCOPE_API_KEY}",
                "Content-Type": "application/json",
            },
            json=payload,
        )
        resp.raise_for_status()
        data = resp.json()

    choices = data.get("choices", [])
    if not choices:
        return ""
    return choices[0].get("message", {}).get("content", "") or ""
